from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper" / "研究进展报告_DREAMER_NRC.md"
OUTPUT = ROOT / "paper" / "研究进展报告_DREAMER_NRC.docx"
MAIN_FIGURE = ROOT / "public_data" / "dreamer_pilot" / "results" / "DREAMER_NRC_直观对照图.png"
WEB_FIGURE = ROOT / "visualization" / "dreamer-vad-3d" / "verification" / "desktop.png"

RED = "B5403E"
INK = "202A35"
MUTED = "607080"
LIGHT = "F2F5F7"
LINE = "D9E0E5"
TEAL = "287B78"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="Microsoft YaHei", size=None, bold=None, color=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^]]+\]\([^)]+\))")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            set_font(paragraph.add_run(text[pos : match.start()]))
        token = match.group(0)
        if token.startswith("**"):
            set_font(paragraph.add_run(token[2:-2]), bold=True)
        elif token.startswith("`"):
            set_font(paragraph.add_run(token[1:-1]), name="Consolas", size=9, color=TEAL)
        else:
            label = token[1 : token.index("]")]
            set_font(paragraph.add_run(label), color=TEAL)
        pos = match.end()
    if pos < len(text):
        set_font(paragraph.add_run(text[pos:]))


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color, before, after in (
        ("Title", 27, INK, 0, 14),
        ("Heading 1", 18, INK, 18, 8),
        ("Heading 2", 13, RED, 12, 5),
        ("Heading 3", 11, TEAL, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Caption CN" not in styles:
        caption = styles.add_style("Caption CN", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = "Microsoft YaHei"
        caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        caption.font.size = Pt(8.5)
        caption.font.color.rgb = RGBColor.from_string(MUTED)
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(8)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("DREAMER × NRC-VAD 研究进展"), size=8.5, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("研究进展报告  |  "), size=8, color=MUTED)
    add_field(footer, "PAGE")


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(68)
    run = p.add_run("研究进展报告  ·  2026.07.24")
    set_font(run, size=10, bold=True, color=RED)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(10)
    set_font(title.add_run("研究进展报告"), size=29, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(30)
    set_font(
        subtitle.add_run(
            "感情語に対する意味的評価と映像によって喚起された情動体験との"
            "VAD次元別対応関係の検討"
        ),
        size=15,
        bold=True,
        color=TEAL,
    )

    line = doc.add_paragraph()
    line.paragraph_format.space_after = Pt(24)
    set_font(line.add_run("研究题目确定  |  主分析与稳健性复核完成"), size=11, color=MUTED)

    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Cm(4)
    table.columns[1].width = Cm(12)
    items = [
        ("核心数据", "DREAMER：23名参与者 × 18段影片 = 414试次"),
        ("对齐对象", "9类诱发情绪 VAD × NRC-VAD v2.1 同名情绪词"),
        ("阶段判断", "足以形成探索性论文，不替代同参与者三层验证实验"),
    ]
    for idx, (key, value) in enumerate(items):
        for cell in table.rows[idx].cells:
            set_cell_margins(cell, 130, 140, 130, 140)
        set_cell_shading(table.cell(idx, 0), RED if idx == 2 else LIGHT)
        set_cell_shading(table.cell(idx, 1), "F8FAFB")
        p1 = table.cell(idx, 0).paragraphs[0]
        p2 = table.cell(idx, 1).paragraphs[0]
        set_font(p1.add_run(key), size=9.5, bold=True, color="FFFFFF" if idx == 2 else INK)
        set_font(p2.add_run(value), size=9.5, bold=idx == 2, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.paragraph_format.space_after = Pt(8)
    set_font(p.add_run("一句话结论"), size=10, bold=True, color=RED)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.right_indent = Cm(0.4)
    p.paragraph_format.line_spacing = 1.4
    set_font(
        p.add_run(
            "情绪词与诱发体验在 Valence 上高度对应，在 Arousal 上存在尺度压缩和排序变化，"
            "在 Dominance 上暂未显示可靠对应；原始距离与校准残差必须同时解释。"
        ),
        size=12,
        bold=True,
        color=INK,
    )
    doc.add_page_break()


def add_markdown_table(doc: Document, rows: list[str]) -> None:
    parsed = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in rows]
    if len(parsed) < 2:
        return
    data = [parsed[0]] + parsed[2:]
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(data):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, INK if i == 0 else (LIGHT if i % 2 == 0 else "FFFFFF"))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value)
            for run in p.runs:
                set_font(run, size=8.3, bold=(i == 0), color="FFFFFF" if i == 0 else INK)
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc: Document, path: Path, caption: str, width: float) -> None:
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph(style="Caption CN")
    cp.add_run(caption)


def render_body(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    start = next(i for i, line in enumerate(lines) if line.strip() == "## 摘要")
    lines = lines[start:]
    i = 0
    forced_breaks = {"## 2.", "## 6.", "## 9.", "## 参考文献"}
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            add_markdown_table(doc, block)
            continue
        if line.startswith("## "):
            if any(line.startswith(prefix) for prefix in forced_breaks) and len(doc.paragraphs) > 2:
                doc.add_page_break()
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, line[3:])
            if line.startswith("## 5."):
                add_figure(doc, MAIN_FIGURE, "图1  NRC 情绪词坐标与 DREAMER 诱发体验均值的直观对照", 6.25)
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[4:])
            i += 1
            continue
        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Cm(0.7)
            add_inline(p, re.sub(r"^\d+\. ", "", line))
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.7)
            add_inline(p, line[2:])
            i += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.right_indent = Cm(0.7)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(8)
            add_inline(p, line[2:])
            for run in p.runs:
                set_font(run, size=11, bold=True, color=TEAL)
            i += 1
            continue
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        add_inline(p, line.rstrip("  "))
        if line.startswith("交互模型位置："):
            add_figure(doc, WEB_FIGURE, "图2  Web 3D 模型：两套 VAD 坐标及情绪差异向量", 6.1)
        i += 1


def main() -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    render_body(doc, SOURCE.read_text(encoding="utf-8"))

    core = doc.core_properties
    core.title = "感情语意义评价与影片诱发情动体验的VAD次元别对应关系研究进展报告"
    core.subject = "DREAMER 与 NRC-VAD 跨语境对齐研究"
    core.author = "研究项目组"
    core.keywords = "VAD, DREAMER, NRC-VAD, 情绪诱发, 词义理解"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print("progress report DOCX generated")


if __name__ == "__main__":
    main()
