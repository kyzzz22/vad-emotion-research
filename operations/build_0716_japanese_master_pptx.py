from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
MASTER = Path(r"E:\dolylab\进展报告\0716\716 - コピー.pptx")
OUTPUT = ROOT / "paper" / "0716_進捗報告_母版準拠_日本語版.pptx"

FIG = {
    "vad_scatter": ROOT / "public_data" / "dreamer_pilot" / "results" / "dreamer_nrc_dimension_scatter.png",
    "vad_distance": ROOT / "public_data" / "dreamer_pilot" / "results" / "dreamer_nrc_vad_distance.png",
    "case_recovery": ROOT / "public_data" / "case_recovery" / "week2" / "overall_recovery_curves_ci.png",
    "case_consistency": ROOT / "public_data" / "case_recovery" / "week2" / "cross_condition_auc_scatter.png",
}

FONT_TITLE = "HGP創英角ｺﾞｼｯｸUB"
FONT_BODY = "HG丸ｺﾞｼｯｸM-PRO"
FONT_BOLD = "HGPSoeiKakugothicUB"
INK = "000000"
BLUE = "1F4E79"
RED = "9E2F2F"
GRAY = "666666"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def remove_all_slides(prs: Presentation) -> None:
    slide_id_list = prs.slides._sldIdLst
    for slide_id in list(slide_id_list):
        prs.part.drop_rel(slide_id.rId)
        slide_id_list.remove(slide_id)


def add_textbox(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    font: str = FONT_BODY,
    size: float = 20,
    bold: bool = False,
    color: str = INK,
    align=PP_ALIGN.LEFT,
    margin: float = 0.0,
    line_spacing: float | None = None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for idx, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing is not None:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def add_bullets(slide, lines: list[str], x: float, y: float, w: float, h: float, *, size: float = 21):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = f"・{line}"
        run.font.name = FONT_BODY
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(INK)
    return box


def add_header_footer(slide, title: str, page: int) -> None:
    # Coordinates are taken from slide 2 of the provided master.
    add_textbox(slide, title, 0.92, 0.31, 11.5, 0.98, font=FONT_TITLE, size=30, bold=False)
    add_textbox(slide, "2026", 0.92, 6.95, 3.0, 0.4, font=FONT_BODY, size=12)
    add_textbox(slide, str(page), 9.42, 6.95, 3.0, 0.4, font=FONT_BODY, size=12, align=PP_ALIGN.RIGHT)


def add_image_contain(slide, image_path: Path, x: float, y: float, w: float, h: float) -> None:
    with Image.open(image_path) as im:
        iw, ih = im.size
    aspect = iw / ih
    box_aspect = w / h
    if aspect >= box_aspect:
        final_w = w
        final_h = w / aspect
        final_x = x
        final_y = y + (h - final_h) / 2
    else:
        final_h = h
        final_w = h * aspect
        final_x = x + (w - final_w) / 2
        final_y = y
    slide.shapes.add_picture(
        str(image_path),
        Inches(final_x),
        Inches(final_y),
        width=Inches(final_w),
        height=Inches(final_h),
    )


def add_two_column(slide, left_title: str, left_lines: list[str], right_title: str, right_lines: list[str]) -> None:
    add_textbox(slide, left_title, 1.05, 1.42, 5.1, 0.34, font=FONT_BOLD, size=20, color=BLUE)
    add_bullets(slide, left_lines, 1.05, 1.9, 5.2, 3.9, size=18.5)
    add_textbox(slide, right_title, 7.0, 1.42, 5.1, 0.34, font=FONT_BOLD, size=20, color=BLUE)
    add_bullets(slide, right_lines, 7.0, 1.9, 5.2, 3.9, size=18.5)


def copy_slide_background(master_slide, target_slide) -> None:
    # Keep non-text decorative elements from the master slide when present.
    for shape in master_slide.shapes:
        if getattr(shape, "has_text_frame", False):
            continue
        target_slide.shapes._spTree.insert_element_before(deepcopy(shape.element), "p:extLst")


def build() -> None:
    prs = Presentation(str(MASTER))
    master_slides = list(prs.slides)
    remove_all_slides(prs)

    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    copy_slide_background(master_slides[0], slide)
    add_textbox(slide, "進捗報告", 1.67, 1.73, 10.0, 1.81, font=FONT_TITLE, size=60, align=PP_ALIGN.CENTER)
    add_textbox(slide, "７/16", 2.73, 4.06, 7.86, 0.5, font=FONT_TITLE, size=24, align=PP_ALIGN.CENTER)
    add_textbox(slide, "MA26006\n姜晰頡", 1.67, 5.58, 10.0, 1.81, font=FONT_BOLD, size=24, align=PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank)
    copy_slide_background(master_slides[1], slide)
    add_header_footer(slide, "本日の構成", 2)
    add_bullets(
        slide,
        [
            "今週の分析進捗と全体判断",
            "主線：DREAMER と NRC-VAD による三次元 VAD 空間の対応",
            "副線：CASE データによる情動変化と生理的回復過程",
            "次の作業：VAD 主線の論文化と副線の位置づけ整理",
        ],
        1.1,
        1.55,
        10.8,
        4.8,
        size=23,
    )

    slide = prs.slides.add_slide(blank)
    copy_slide_background(master_slides[1], slide)
    add_header_footer(slide, "今週の全体進捗", 3)
    add_two_column(
        slide,
        "主線：VAD 方向",
        [
            "DREAMER の 414 試行を解析し、NRC-VAD と同一尺度に変換した",
            "9 種類の情動語について、次元相関・距離・校正残差を算出した",
            "Valence は安定した対応を示し、Dominance は不安定であった",
        ],
        "副線：情動変化方向",
        [
            "CASE の 30 名全サンプルで回復曲線を作成した",
            "EDA/SCL では scary 条件後の残留反応が明確であった",
            "副線は主線を補足する回復動態の分析として整理する",
        ],
    )

    slide = prs.slides.add_slide(blank)
    copy_slide_background(master_slides[1], slide)
    add_header_footer(slide, "VAD 方向：研究目的とデータ", 4)
    add_textbox(
        slide,
        "研究目的",
        1.0,
        1.35,
        2.5,
        0.35,
        font=FONT_BOLD,
        size=20,
        color=BLUE,
    )
    add_textbox(
        slide,
        "情動語の意味的 VAD と、映像刺激後の主観的 VAD が、\nカテゴリーレベルの三次元空間でどの程度対応するかを検討する。",
        1.0,
        1.85,
        11.0,
        0.95,
        font=FONT_BODY,
        size=22,
        line_spacing=1.15,
    )
    add_bullets(
        slide,
        [
            "DREAMER：23 名、18 本の映像、414 試行、Valence・Arousal・Dominance 評定",
            "NRC-VAD v2.1：9 種類の同名情動語を使用",
            "両データを [-1, 1] に線形変換し、同一座標系で比較した",
            "本研究は個人内予測ではなく、カテゴリーレベルの探索的対応分析である",
        ],
        1.0,
        3.25,
        11.2,
        2.7,
        size=19.5,
    )

    slide = prs.slides.add_slide(blank)
    copy_slide_background(master_slides[1], slide)
    add_header_footer(slide, "VAD 方向：次元別の対応", 5)
    add_image_contain(slide, FIG["vad_scatter"], 0.75, 1.25, 11.85, 4.18)
    add_textbox(
        slide,
        "Valence は強く安定した対応を示した。一方、Arousal は線形対応があるが順位は不安定であり、Dominance は明確な対応を示さなかった。",
        0.92,
        5.72,
        11.4,
        0.75,
        font=FONT_BODY,
        size=17,
    )

    slide = prs.slides.add_slide(blank)
    copy_slide_background(master_slides[1], slide)
    add_header_footer(slide, "VAD 方向：具体的な結果", 6)
    add_textbox(slide, "次元ごとの結論", 1.0, 1.33, 3.0, 0.35, font=FONT_BOLD, size=20, color=BLUE)
    add_bullets(
        slide,
        [
            "Valence：Pearson r = .875, p = .0020；Spearman も有意であり、最も安定した対応であった",
            "Arousal：Pearson r = .832, p = .0054；線形対応は強いが、Spearman は有意でなく、情動間の順位は不安定であった",
            "Dominance：Pearson r = -.121, p = .757；語彙規範と誘発経験の間に明確な対応は見られなかった",
        ],
        1.0,
        1.82,
        11.2,
        2.25,
        size=18.2,
    )
    add_textbox(slide, "解釈", 1.0, 4.45, 2.0, 0.35, font=FONT_BOLD, size=20, color=BLUE)
    add_bullets(
        slide,
        [
            "語彙的な情動理解と映像後の主観経験は、快・不快の軸ではかなり近い構造を持つ",
            "覚醒度は全体として同じ方向に動くが、どの情動がより高覚醒かという順位は文脈に依存する",
            "Dominance は、評定方向・課題理解・概念そのものの違いを再確認すべき次元である",
        ],
        1.0,
        4.95,
        11.2,
        1.15,
        size=17.2,
    )

    slide = prs.slides.add_slide(blank)
    copy_slide_background(master_slides[1], slide)
    add_header_footer(slide, "VAD 方向：距離と校正残差", 7)
    add_image_contain(slide, FIG["vad_distance"], 1.0, 1.25, 7.15, 4.85)
    add_bullets(
        slide,
        [
            "原距離は、二つの座標が幾何学的にどれだけ離れているかを示す",
            "happiness は近く、sadness は遠い",
            "LOEO 校正残差は、全体的な尺度差を除いた情動固有のずれを示す",
            "calmness は校正後も大きなずれを示した",
        ],
        8.55,
        1.55,
        3.85,
        4.6,
        size=15.8,
    )

    slide = prs.slides.add_slide(blank)
    copy_slide_background(master_slides[1], slide)
    add_header_footer(slide, "VAD 方向：情動別の解釈", 8)
    add_two_column(
        slide,
        "原距離から見えること",
        [
            "happiness は最も近く、語彙的意味と誘発経験が比較的一致していた",
            "sadness は最も遠く、語彙規範の極端さと映像後評定の平均化がずれを生んだ可能性がある",
            "fear・disgust・surprise も原距離は大きく、刺激文脈による経験の圧縮が考えられる",
        ],
        "校正残差から見えること",
        [
            "fear は原距離が大きいが、校正後の残差は小さく、主に全体的な尺度差で説明できる",
            "calmness は校正後も残差が最大であり、情動固有のずれとして議論しやすい",
            "excitement・anger は校正後の残差が小さく、VAD 空間内で比較的再現しやすい情動である",
        ],
    )

    slide = prs.slides.add_slide(blank)
    copy_slide_background(master_slides[1], slide)
    add_header_footer(slide, "VAD 方向：現時点の解釈", 9)
    add_two_column(
        slide,
        "言えること",
        [
            "情動語と誘発経験は、Valence において共通した構造を持つ",
            "Arousal では対応と再順位化が同時に見られる",
            "距離と校正残差を分けることで、尺度差と情動固有差を区別できる",
        ],
        "言えないこと",
        [
            "同一個人の語義判断が誘発経験を予測するとは言えない",
            "意味的 VAD と誘発 VAD が同一測定であるとは言えない",
            "Dominance の不対応を、直ちに生理的基盤の欠如とは解釈できない",
        ],
    )

    slide = prs.slides.add_slide(blank)
    copy_slide_background(master_slides[1], slide)
    add_header_footer(slide, "副線：情動変化と回復過程", 10)
    add_image_contain(slide, FIG["case_recovery"], 0.75, 1.25, 11.85, 4.4)
    add_textbox(
        slide,
        "CASE データでは、scary 条件後の SCL 残留反応が amusing 条件より大きく、120 秒の回復区間で差が緩やかに縮小した。",
        0.92,
        5.88,
        11.4,
        0.55,
        font=FONT_BODY,
        size=17,
    )

    slide = prs.slides.add_slide(blank)
    copy_slide_background(master_slides[1], slide)
    add_header_footer(slide, "副線：完了した分析と限界", 11)
    add_image_contain(slide, FIG["case_consistency"], 0.95, 1.35, 5.8, 2.58)
    add_bullets(
        slide,
        [
            "30 名全サンプル、240 trial、2880 個の 10 秒回復 bin を作成した",
            "ECG/EDA 品質確認、混合モデル、感度分析を完了した",
            "SCL の回復 AUC と回復傾きには探索的な個人内一貫性が見られた",
            "ただし各条件 2 trial であるため、安定した trait 指標とは解釈しない",
        ],
        7.05,
        1.35,
        5.15,
        4.3,
        size=17.2,
    )
    add_textbox(
        slide,
        "副線の役割：主線とは別に、情動誘発後の生理的回復動態を扱う補助的研究として位置づける。",
        1.0,
        5.35,
        11.4,
        0.58,
        font=FONT_BODY,
        size=17,
        color=RED,
    )

    slide = prs.slides.add_slide(blank)
    copy_slide_background(master_slides[1], slide)
    add_header_footer(slide, "次の作業", 12)
    add_two_column(
        slide,
        "VAD 主線",
        [
            "Dominance の評定方向を原資料で再確認する",
            "置換検定、bootstrap、1 情動除外感度分析を追加する",
            "名詞マッピングを主分析、形容詞マッピングを感度分析として整理する",
            "方法・結果・考察を論文形式にまとめる",
        ],
        "情動変化副線",
        [
            "CASE 第二週の結果を短い補助資料として整理する",
            "SCL を中心に、回復動態の図とモデル結果を明確化する",
            "主観的 arousal と生理的回復を同一概念として扱わない",
            "今月の主成果は VAD 論文に集中する",
        ],
    )

    slide = prs.slides.add_slide(blank)
    copy_slide_background(master_slides[1], slide)
    add_header_footer(slide, "まとめ", 13)
    add_textbox(
        slide,
        "今週の結論",
        1.0,
        1.35,
        3.0,
        0.35,
        font=FONT_BOLD,
        size=22,
        color=BLUE,
    )
    add_textbox(
        slide,
        "VAD 方向は、探索的論文としてまとめられる段階まで進んだ。\n最も安定した結果は Valence の対応であり、Arousal は部分的対応、Dominance は不安定であった。\n情動変化方向は、副線として生理的回復動態と個人差の分析を支える。",
        1.0,
        2.0,
        11.1,
        2.0,
        font=FONT_BODY,
        size=23,
        line_spacing=1.1,
    )
    add_textbox(
        slide,
        "次回までに、VAD 主線の感度分析と論文構成を優先して進める。",
        1.0,
        5.0,
        11.1,
        0.5,
        font=FONT_BOLD,
        size=21,
        color=RED,
    )

    prs.save(str(OUTPUT))
    print("Japanese master-based deck generated")


if __name__ == "__main__":
    build()
