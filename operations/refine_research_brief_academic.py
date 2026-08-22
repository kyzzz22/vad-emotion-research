"""Create an academically aligned revision of the two-page research brief.

The script preserves the established A4/two-column layout, figure, table,
affiliation line, and reference list.  It tightens the inferential claims and
uses one terminology system throughout:

- 意味的評価 / 情動体験評価
- 線形対応 / 順位対応 / 値の広がり
- 対応 (association) / 数値的一致 (numerical agreement)
"""

from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper" / "研究概要書_JIANGXIJIE_MA26006_審査指摘反映版.docx"
OUTPUT = ROOT / "paper" / "研究概要書_JIANGXIJIE_MA26006_学術整合版.docx"

FONT_BODY = "ＭＳ 明朝"
FONT_HEAD = "ＭＳ ゴシック"
BODY_SIZE = Pt(10)
REF_SIZE = Pt(9)
CAPTION_SIZE = Pt(10)


def paragraph_starting(document: Document, prefix: str):
    matches = [p for p in document.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one paragraph starting with {prefix!r}, got {len(matches)}"
        )
    return matches[0]


def clear_runs(paragraph) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def set_run_font(run, family: str, size, *, bold: bool = False) -> None:
    run.font.name = family
    run.font.size = size
    run.font.bold = bold
    run.font.italic = False
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), family)


def replace_paragraph(
    paragraph,
    text: str,
    *,
    family: str = FONT_BODY,
    size=BODY_SIZE,
    bold: bool = False,
) -> None:
    first_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            first_rpr = copy.deepcopy(run._r.rPr)
            break
    clear_runs(paragraph)
    run = paragraph.add_run(text)
    if first_rpr is not None:
        run._r.insert(0, first_rpr)
    set_run_font(run, family, size, bold=bold)


def replace_inline_heading(paragraph, heading: str, body: str) -> None:
    clear_runs(paragraph)
    heading_run = paragraph.add_run(heading)
    set_run_font(heading_run, FONT_HEAD, BODY_SIZE, bold=True)
    body_run = paragraph.add_run(body)
    set_run_font(body_run, FONT_BODY, BODY_SIZE)


def replace_table_cell(cell, text: str, *, header: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    replace_paragraph(
        paragraph,
        text,
        family=FONT_HEAD if header else FONT_BODY,
        size=BODY_SIZE,
        bold=header,
    )


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    document = Document(SOURCE)

    replace_paragraph(
        paragraph_starting(document, "指導教員"),
        "指導教員　菅谷みどり\tMA26006　姜晰頡",
    )

    replace_paragraph(
        paragraph_starting(document, "感情計算では"),
        "感情計算では，異なる対象の感情を共通の数値で扱う必要がある．そこで，感情価"
        "（Valence，V），覚醒度（Arousal，A），支配性（Dominance，D）の3次元からなる"
        "VADモデルが用いられる．Vは快・不快，Aは活性度，Dは制御感を表す[1]．",
    )
    replace_paragraph(
        paragraph_starting(document, "VADを用いれば"),
        "VADを用いると，感情語の意味と映像後の情動体験を同じ座標空間に配置できる．"
        "そのため，辞書VADを情動体験の参照値に用いることが考えられる．しかし，前者は"
        "語の一般的意味，後者は特定状況で生じた体験を対象とする[5，6]．以下，前者を"
        "「意味的評価」，後者を「情動体験評価」と呼ぶ．評価対象が異なるため，辞書値の"
        "直接利用には妥当性の問題が残る．",
    )
    replace_paragraph(
        paragraph_starting(document, "この問題を検証するには"),
        "この妥当性を検討するには，両評価の直接比較が必要である．NRC-VADは英語語彙の"
        "VAD規範値，DREAMERは映像視聴後のVAD評価，CASEは映像視聴中の連続VA評価を"
        "収録する[2-4]．ただし，参加者，刺激，課題は異なる．したがって，同じ感情"
        "カテゴリー名を介した対応がV・A・Dで同程度に保たれるかは明らかでない．",
    )
    replace_paragraph(
        paragraph_starting(document, "また，高い相関があっても"),
        "高い相関は数値的一致を意味しない．相対方向が同じでも，情動体験評価が中立点"
        "付近に集中すれば，辞書値を代替できない．そこで，線形対応，順位対応，値の"
        "広がりを分けて検討する．本研究はこの比較枠組みを示し，次元別の対応特性を"
        "探索する．これにより，辞書値を参照値として利用する際の基礎的判断材料を提供する．",
    )

    replace_paragraph(
        paragraph_starting(document, "以上から，NRC-VAD"),
        "NRC-VADとDREAMERに共通する9感情カテゴリーを対象とする．意味的評価と情動体験"
        "評価の①線形対応，②順位対応，③値の広がりをV・A・D別に検討する．目的は同一性の"
        "証明ではなく，次元ごとの対応特性と利用上の限界を明らかにすることである．",
    )
    replace_paragraph(
        paragraph_starting(document, "共通する感情カテゴリー名"),
        "共通する感情カテゴリー名を対応づけ，各尺度の理論上の端点を[-1，1]に線形変換する．"
        "この変換は測定特性の同一性を保証しない．そこで，Pearson相関で線形対応，"
        "Spearman相関で順位対応，回帰傾きで値の広がりを評価する．これにより，"
        "「対応」と「数値的一致」を区別する．",
    )

    p41 = paragraph_starting(document, "4.1 評価目的")
    replace_inline_heading(
        p41,
        "4.1 評価目的　",
        "比較枠組みを9感情カテゴリーに適用する．主要な推論にはVADを含むDREAMERを用いる．"
        "CASEはVAのみで共通カテゴリーも4つのため，連続評定への手順確認に限定する．",
    )
    p42 = paragraph_starting(document, "4.2 評価環境")
    replace_inline_heading(
        p42,
        "4.2 評価環境　",
        "既存データを二次分析する．NRC-VADでは，DREAMERのラベルと同名の英語名詞1語を"
        "参照値とした[3]．これは辞書全体ではなく，カテゴリーラベル単位の操作化である．"
        "DREAMERの23名×18映像＝414試行を9カテゴリーに集約した[2]．各カテゴリーは2映像を"
        "含む．CASEは30名×8映像の連続VA評価を含む[4]．",
    )
    replace_paragraph(
        paragraph_starting(document, "試行を9カテゴリーに集約し"),
        "試行を9カテゴリーに集約した．Pearson相関で線形対応，Spearman相関で順位対応を"
        "評価した．推論単位が9と少ないため，全9!通りの正確置換検定を行った．各相関系列の"
        "V・A・DをHolm法で補正した．値の広がりは回帰傾きで近似した．参加者23名と各"
        "カテゴリー内の2映像を復元抽出する5,000回の階層bootstrapで95%信頼区間を求めた．"
        "語形置換，1カテゴリー除外，Dの符号反転による感度分析も行った．",
    )

    replace_paragraph(
        paragraph_starting(document, "以上の3段階で評価した結果"),
        "Vでは線形対応と順位対応がともに支持された．Aでは線形対応のみが支持された．"
        "Dではいずれも支持されず，bootstrap区間も0を含んだ（図１，表１）．したがって，"
        "対応特性はV・A・Dで異なった．",
    )
    replace_paragraph(
        paragraph_starting(document, "線形対応が得られたVとA"),
        "VとAの回帰傾きはそれぞれ.520，.341で，いずれも1未満であった（表１）．"
        "情動体験評価は中立点付近に集中し，高相関でも数値は一致しない．ただし，傾きは"
        "値の広がりの近似であり，尺度，刺激，標本の影響を分離できない．形容詞への置換と"
        "Dの符号反転で主要な判断は変わらなかった．一方，Aはcalmnessを除くとr=.482まで"
        "低下した．CASEにも同じ手順を適用できたが，4カテゴリーのため外的再現とはみなさない．",
    )
    replace_paragraph(
        paragraph_starting(document, "ValenceではPearson相関"),
        "Vでは線形対応と順位対応が支持された．1カテゴリー除外後のrも.848～.922であった．"
        "したがって，相対構造は比較的安定していた．ただし，傾きが.520であるため，"
        "意味的評価は情動体験評価の方向を捉えても，数値を代替できない．",
    )
    replace_paragraph(
        paragraph_starting(document, "ArousalではPearson相関"),
        "Aでは線形対応が支持された一方，順位対応は支持されなかった．calmnessを除くと"
        "rも大きく低下した．したがって，線形対応はカテゴリー構成に依存する可能性があり，"
        "安定した順位構造の証拠は不十分である．",
    )
    replace_paragraph(
        paragraph_starting(document, "Dominanceでは線形対応も順位対応も"),
        "Dでは線形対応も順位対応も支持されなかった．意味的評価のDは語の一般的支配感，"
        "情動体験評価のDは場面内の主体性や状況統制にも左右されうる[7]．ただし，本分析は"
        "この機序を検証していない．したがって，Dの対応が得られなかった理由を構念差だけでは"
        "説明できない．",
    )

    replace_paragraph(
        paragraph_starting(document, "以上から，意味的評価と情動体験評価"),
        "探索的分析の結果，対応特性はV・A・Dで異なった．Vの相対構造は比較的安定したが，"
        "値は縮小した．Aは線形対応に限られ，カテゴリー構成にも敏感であった．Dでは対応の"
        "証拠が得られなかった．したがって，辞書VADは情動体験評価の代替値ではなく，利用時は"
        "次元別の検証と較正が必要である．対応と数値的一致を区別する枠組みと，次元別の"
        "探索的証拠を示した点が本研究の貢献である．",
    )
    replace_paragraph(
        paragraph_starting(document, "本分析は9カテゴリー"),
        "推論単位は9カテゴリーであり，各カテゴリーは名詞1語と2映像に限られる．したがって，"
        "辞書全体や多様な刺激へ一般化できない．また，データセット間で参加者と課題が異なる"
        "ため，評価対象の効果を分離できず，個人内対応も判断できない．CASEも外的再現の"
        "証拠にはならない．",
    )
    replace_paragraph(
        paragraph_starting(document, "まず，感情数・映像数・類義語"),
        "まず，複数の類義語，感情カテゴリー，映像を用いて外部データで再検証する．次に，"
        "同一参加者が感情語の意味と映像後体験を同一尺度で評価する実験を行う．刺激を"
        "交差させた混合効果モデルにより，個人内対応と較正式の一般化可能性を検討する．",
    )

    replace_paragraph(
        paragraph_starting(document, "図1"),
        "図１　意味的評価と情動体験評価の次元別対応（9感情カテゴリー）",
        family=FONT_HEAD,
        size=CAPTION_SIZE,
    )
    replace_paragraph(
        paragraph_starting(document, "表1"),
        "表１　VAD次元別の対応指標（9感情カテゴリー）",
        family=FONT_HEAD,
        size=CAPTION_SIZE,
    )

    table = document.tables[0]
    replace_table_cell(table.cell(4, 0), "値の広がりの近似\n回帰傾き")

    # Normalize the reference typography and the table after the local edits.
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("["):
            for run in paragraph.runs:
                set_run_font(run, FONT_BODY, REF_SIZE)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        FONT_HEAD if row_index == 0 else FONT_BODY,
                        BODY_SIZE,
                        bold=row_index == 0,
                    )

    # Mirror the existing picture descriptions to Word's accessibility field.
    # The retained template stores descriptions in pic:cNvPr, while the audit
    # and Word's alt-text UI read wp:docPr.
    for shape in document.inline_shapes:
        picture_props = shape._inline.graphic.graphicData.pic.nvPicPr.cNvPr
        description = picture_props.get("descr")
        if description:
            shape._inline.docPr.set("descr", description)

    document.core_properties.subject = (
        "意味的評価と情動体験評価のVAD次元別対応を検討する二頁研究概要書"
    )
    document.core_properties.comments = (
        "論理連鎖，統計的主張，限界，貢献を再点検し，意味的評価・情動体験評価・"
        "線形対応・順位対応・値の広がりの用語を統一した学術整合版．"
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
