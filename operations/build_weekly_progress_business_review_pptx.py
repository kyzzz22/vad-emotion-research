from __future__ import annotations

import shutil
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(
    r"C:\Users\1\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.0"
    r"\skills\artifact-template-business-review\assets\reference.pptx"
)
OUTPUT = ROOT / "paper" / "本周研究进展汇报_business_review模板_VAD主线_情绪变化副线.pptx"

IMG = {
    "vad_overview": ROOT / "public_data" / "dreamer_pilot" / "results" / "DREAMER_NRC_直观对照图.png",
    "vad_scatter": ROOT / "public_data" / "dreamer_pilot" / "results" / "dreamer_nrc_dimension_scatter.png",
    "vad_distance": ROOT / "public_data" / "dreamer_pilot" / "results" / "dreamer_nrc_vad_distance.png",
    "vad_3d": ROOT / "visualization" / "dreamer-vad-3d" / "verification" / "desktop.png",
    "recovery": ROOT / "public_data" / "case_recovery" / "week2" / "overall_recovery_curves_ci.png",
    "consistency": ROOT / "public_data" / "case_recovery" / "week2" / "cross_condition_auc_scatter.png",
}

BLUE = "3157B7"
INK = "111827"
MUTED = "667085"
PALE = "F5F7FA"
LINE = "D0D5DD"
RED = "B5403E"
TEAL = "287B78"
GREEN = "4D7C59"
GOLD = "A66F16"
WHITE = "FFFFFF"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def remove_all_slides(prs: Presentation) -> None:
    slide_id_list = prs.slides._sldIdLst
    for slide_id in list(slide_id_list):
        prs.part.drop_rel(slide_id.rId)
        slide_id_list.remove(slide_id)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 14,
    color: str = INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    font: str = "Microsoft YaHei",
    margin: float = 0.02,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_multiline(
    slide,
    lines: list[str],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 13,
    color: str = INK,
    bullet: bool = False,
    bold_first: bool = False,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        if bullet:
            p.text = f"• {line}"
        else:
            p.text = line
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.font.bold = bool(bold_first and i == 0)
    return box


def add_rule(slide, y: float = 1.15) -> None:
    slide.shapes.add_connector(
        1, Inches(0), Inches(y), Inches(13.333), Inches(y)
    ).line.color.rgb = rgb(INK)


def add_footer(slide, number: int) -> None:
    add_text(slide, str(number), 0.45, 6.95, 0.3, 0.18, size=8, color=MUTED, bold=True)
    add_text(slide, "Weekly research progress", 10.2, 6.95, 2.5, 0.18, size=7.5, color=MUTED, align=PP_ALIGN.RIGHT)


def add_header(slide, title: str, kicker: str, number: int) -> None:
    add_text(slide, kicker, 0.62, 0.38, 4.0, 0.25, size=8.5, color=MUTED, bold=True)
    add_text(slide, title, 0.62, 0.68, 10.8, 0.42, size=22, color=BLUE, bold=True)
    add_rule(slide)
    add_footer(slide, number)


def add_card(slide, title: str, body: str, x: float, y: float, w: float, h: float, accent: str = BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(WHITE)
    shape.line.color.rgb = rgb(LINE)
    shape.line.width = Pt(0.75)
    add_text(slide, title, x + 0.18, y + 0.16, w - 0.36, 0.28, size=11.5, color=accent, bold=True)
    add_multiline(slide, [body], x + 0.18, y + 0.55, w - 0.36, h - 0.68, size=10.2, color=INK)


def add_metric(slide, value: str, label: str, x: float, y: float, w: float, color: str):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.82))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(PALE)
    shape.line.color.rgb = rgb(LINE)
    add_text(slide, value, x + 0.08, y + 0.12, w - 0.16, 0.25, size=18, color=color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, label, x + 0.08, y + 0.49, w - 0.16, 0.2, size=8.4, color=MUTED, align=PP_ALIGN.CENTER)


def add_image_contain(slide, image_path: Path, x: float, y: float, w: float, h: float, border: bool = True):
    with Image.open(image_path) as im:
        iw, ih = im.size
    image_aspect = iw / ih
    box_aspect = w / h
    if image_aspect >= box_aspect:
        final_w = w
        final_h = w / image_aspect
        final_x = x
        final_y = y + (h - final_h) / 2
    else:
        final_h = h
        final_w = h * image_aspect
        final_x = x + (w - final_w) / 2
        final_y = y
    if border:
        rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        rect.fill.solid()
        rect.fill.fore_color.rgb = rgb(WHITE)
        rect.line.color.rgb = rgb(LINE)
        rect.line.width = Pt(0.75)
    slide.shapes.add_picture(str(image_path), Inches(final_x), Inches(final_y), width=Inches(final_w), height=Inches(final_h))


def add_table_like(slide, rows: list[tuple[str, str, str]], x: float, y: float, w: float, row_h: float = 0.44):
    col_w = [w * 0.24, w * 0.31, w * 0.45]
    headers = ("维度/主题", "本周证据", "讲法")
    cur_y = y
    for r, row in enumerate([headers, *rows]):
        cur_x = x
        for c, value in enumerate(row):
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cur_x), Inches(cur_y), Inches(col_w[c]), Inches(row_h))
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(BLUE if r == 0 else (PALE if r % 2 == 0 else WHITE))
            shape.line.color.rgb = rgb(LINE)
            add_text(
                slide,
                value,
                cur_x + 0.07,
                cur_y + 0.08,
                col_w[c] - 0.14,
                row_h - 0.14,
                size=8.6 if r else 8.8,
                color=WHITE if r == 0 else INK,
                bold=r == 0,
            )
            cur_x += col_w[c]
        cur_y += row_h


def add_title_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[7])
    add_text(slide, "dolylab research update", 0.62, 0.46, 3.4, 0.24, size=9.5, color=MUTED, bold=True)
    add_text(slide, "本周研究进展汇报", 0.62, 0.92, 6.6, 0.5, size=25, color=BLUE, bold=True)
    add_text(slide, "VAD 主线与情绪变化副线", 0.62, 1.48, 6.6, 0.46, size=23, color=BLUE, bold=True)
    add_rule(slide, 2.35)
    add_text(slide, "2026.07.15", 0.62, 3.16, 1.9, 0.24, size=10.5, color=BLUE, bold=True)
    add_multiline(
        slide,
        [
            "主线：DREAMER × NRC-VAD 三维情绪空间对齐",
            "副线：CASE 情绪变化与生理恢复动态",
            "目标：展示本周完成的分析、阶段判断和下一步任务",
        ],
        0.62,
        3.62,
        5.3,
        1.0,
        size=11.2,
        color=INK,
    )
    add_image_contain(slide, IMG["vad_overview"], 7.35, 0.55, 5.25, 5.85, border=False)
    add_footer(slide, 1)


def build() -> None:
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(TEMPLATE, OUTPUT)

    prs = Presentation(str(OUTPUT))
    remove_all_slides(prs)

    add_title_slide(prs)

    slide = prs.slides.add_slide(prs.slide_layouts[7])
    add_header(slide, "Agenda", "MEETING FLOW", 2)
    items = [
        ("01", "本周总体进展", "两条方向的完成度和阶段判断"),
        ("02", "VAD 主线", "数据、方法、三维结果和解释边界"),
        ("03", "情绪变化副线", "CASE recovery 分析与实验转化"),
        ("04", "下一步任务", "主线收稿化，副线流程化"),
    ]
    y = 1.75
    for idx, title, body in items:
        add_text(slide, idx, 0.9, y, 0.55, 0.3, size=16, color=BLUE, bold=True)
        add_text(slide, title, 1.65, y, 3.2, 0.28, size=15, color=INK, bold=True)
        add_text(slide, body, 5.05, y + 0.03, 5.8, 0.22, size=10.5, color=MUTED)
        y += 0.82

    slide = prs.slides.add_slide(prs.slide_layouts[7])
    add_header(slide, "Executive Snapshot", "WEEKLY STATUS", 3)
    add_metric(slide, "VAD", "本周展示重点", 0.75, 1.45, 1.8, RED)
    add_metric(slide, "414", "DREAMER 有效试次", 2.85, 1.45, 1.8, TEAL)
    add_metric(slide, "30", "CASE 全样本参与者", 4.95, 1.45, 1.8, BLUE)
    add_metric(slide, "2", "研究方向同步推进", 7.05, 1.45, 1.8, GREEN)
    add_card(slide, "主线判断", "VAD 方向已从方案验证推进到可写论文的探索性结果：问题、数据、图表、局限和下一步分析均已成型。", 0.75, 2.85, 3.55, 1.35, RED)
    add_card(slide, "副线判断", "CASE recovery 分析完成第二周闭环，可作为情绪变化方向中生理恢复动态与个体差异分析的证据支撑。", 4.65, 2.85, 3.55, 1.35, BLUE)
    add_card(slide, "策略判断", "近期不宜继续扩大战线：优先把 VAD 主线写完整，同时把副线实验流程锁定到 pilot 水平。", 8.55, 2.85, 3.55, 1.35, TEAL)
    add_multiline(
        slide,
        ["一句话结论：VAD 主线已经具备探索性论文雏形；情绪变化副线为后续实验提供 recovery 动态分析框架。"],
        0.9,
        5.35,
        11.4,
        0.5,
        size=15,
        color=INK,
        bold_first=True,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[7])
    add_header(slide, "VAD Direction: Research Question and Pipeline", "MAIN LINE", 4)
    add_text(slide, "核心问题", 0.78, 1.45, 1.5, 0.24, size=10, color=RED, bold=True)
    add_text(
        slide,
        "情绪词的语义 VAD 与影片诱发体验的 VAD，\n在类别级三维空间中是否共享结构？",
        0.78,
        1.82,
        5.5,
        0.86,
        size=20,
        color=INK,
        bold=True,
    )
    add_multiline(
        slide,
        [
            "DREAMER：23人 × 18段影片 = 414个 VAD 评分试次",
            "NRC-VAD v2.1：9个同名情绪词，名词主分析，形容词敏感性分析",
            "统一量尺至 [-1, 1]，比较维度相关、原始距离和 LOEO 校准残差",
        ],
        0.78,
        3.0,
        5.8,
        1.35,
        size=12.3,
        bullet=True,
    )
    add_image_contain(slide, IMG["vad_3d"], 7.02, 1.44, 5.45, 4.55)
    add_text(slide, "定位：跨样本、跨任务、类别级探索性对齐，不写成个体内等价测量验证。", 0.78, 6.2, 11.7, 0.28, size=13, color=RED, bold=True)

    slide = prs.slides.add_slide(prs.slide_layouts[7])
    add_header(slide, "VAD Results: Full Overview", "MAIN FIGURE", 5)
    add_image_contain(slide, IMG["vad_overview"], 0.55, 1.34, 12.25, 5.25)
    add_text(slide, "这页只放总图：用于讲清“两个坐标系如何在 V/A/D 上靠近或偏离”。", 0.72, 6.68, 11.8, 0.22, size=9.4, color=MUTED, align=PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(prs.slide_layouts[7])
    add_header(slide, "VAD Results: Dimension-Specific Alignment", "RESULT 1", 6)
    add_image_contain(slide, IMG["vad_scatter"], 0.55, 1.33, 12.25, 3.95)
    add_table_like(
        slide,
        [
            ("Valence", "Pearson r=.875, p=.0020", "强且稳定的跨语境对应"),
            ("Arousal", "Pearson r=.832, p=.0054", "线性对应存在，但排序不稳定"),
            ("Dominance", "Pearson r=-.121, p=.7573", "暂未形成可靠对应，需核查方向与构念"),
        ],
        0.72,
        5.58,
        11.9,
        row_h=0.35,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[7])
    add_header(slide, "VAD Results: Distance and Residuals", "RESULT 2", 7)
    add_image_contain(slide, IMG["vad_distance"], 0.72, 1.35, 5.65, 4.95)
    add_text(slide, "怎么解释这张图", 6.85, 1.46, 2.1, 0.24, size=10, color=RED, bold=True)
    add_multiline(
        slide,
        [
            "原始距离回答：两套 VAD 坐标在几何上离多远。",
            "happiness 最近，sadness 最远，但原始距离混合了全局尺度压缩。",
            "LOEO 校准回答：扣除整体线性尺度后，哪些情绪仍有特异性偏离。",
            "校准后 excitement、fear、anger 残差较小，calmness 残差最大。",
        ],
        6.85,
        1.86,
        5.55,
        2.3,
        size=12.2,
        bullet=True,
    )
    add_card(slide, "论文贡献", "同时报告“原始距离”和“校准残差”，避免把单一重合度排名误写成心理等价性。", 6.85, 4.65, 5.35, 1.0, TEAL)

    slide = prs.slides.add_slide(prs.slide_layouts[7])
    add_header(slide, "Interpretation Boundary", "WHAT WE CAN / CANNOT SAY", 8)
    add_text(slide, "当前证据支持", 0.9, 1.45, 2.2, 0.24, size=11, color=GREEN, bold=True)
    add_multiline(
        slide,
        [
            "情绪词规范与影片诱发体验在 Valence 上共享明显类别结构。",
            "Arousal 有整体线性对应，但存在尺度压缩和类别重排。",
            "Dominance 是当前最不稳定维度，需要作为讨论和敏感性重点。",
            "跨语境 VAD 对齐框架可服务后续同参与者三层实验。",
        ],
        0.9,
        1.9,
        5.35,
        2.6,
        size=12.5,
        bullet=True,
    )
    add_text(slide, "当前不能声称", 7.0, 1.45, 2.2, 0.24, size=11, color=RED, bold=True)
    add_multiline(
        slide,
        [
            "不能证明同一个人的词义判断可以预测其诱发体验。",
            "不能把语义 VAD 与诱发 VAD 写成同一心理过程。",
            "不能把核心类别相关的样本量写成 414；核心单位是 9 类情绪。",
            "不能把生理模型未显著解释为“生理信号无效”。",
        ],
        7.0,
        1.9,
        5.35,
        2.6,
        size=12.5,
        bullet=True,
    )
    add_text(
        slide,
        "建议题目：From Emotion Words to Elicited Experience: Dimension-Specific Alignment Between NRC-VAD and DREAMER",
        0.9,
        5.72,
        11.5,
        0.33,
        size=12.2,
        color=BLUE,
        bold=True,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[7])
    add_header(slide, "Side Line: CASE Recovery Dynamics", "EMOTION CHANGE", 9)
    add_image_contain(slide, IMG["recovery"], 0.55, 1.33, 12.25, 4.25)
    add_text(slide, "主要发现：scary 后 SCL 残留显著高于 amusing，并在 120 秒内缓慢缩小；HR 有平均条件差异，但恢复形状证据较弱。", 0.72, 6.03, 11.8, 0.34, size=12.3, color=INK, bold=True)

    slide = prs.slides.add_slide(prs.slide_layouts[7])
    add_header(slide, "CASE Week 2: Completed Analyses and Limits", "SIDE LINE STATUS", 10)
    add_card(slide, "完成情况", "30人全样本预处理；240 trial；2880个10秒恢复箱；ECG/EDA 质控；混合模型；敏感性分析；跨条件个体一致性。", 0.75, 1.45, 3.7, 1.35, BLUE)
    add_card(slide, "可以写入", "SCL 条件差异和恢复缩小趋势较稳健；HR 可报告平均差异，但不强调恢复形状。", 4.8, 1.45, 3.7, 1.35, TEAL)
    add_card(slide, "需要保守", "主观 arousal 不能通过全部敏感性检验；结果限定于 CASE 的 amusing/scary，不推广到所有情绪。", 8.85, 1.45, 3.7, 1.35, RED)
    add_image_contain(slide, IMG["consistency"], 0.95, 3.25, 5.65, 2.65)
    add_multiline(
        slide,
        [
            "探索性个体一致性：SCL 绝对恢复 AUC 与恢复斜率有中等正相关。",
            "每个条件只有2个 trial，因此应写成“初步证据”，而不是稳定 trait 测量。",
            "这条副线的价值是提供 recovery 动态特征和质控分析模板。",
        ],
        7.05,
        3.45,
        5.15,
        1.7,
        size=12.1,
        bullet=True,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[7])
    add_header(slide, "How the Side Line Supports the Emotion-Change Direction", "SIDE LINE TRANSFER", 11)
    add_text(slide, "情绪诱发后的多模态生理恢复过程与个体差异量化", 0.75, 1.42, 10.8, 0.35, size=17, color=INK, bold=True)
    steps = [
        ("对象", "正性/负性情绪诱发后的恢复过程"),
        ("窗口", "刺激结束后 recovery 动态"),
        ("指标", "HR/HRV、EDA/SCL\n可扩展 EEG 指标"),
        ("特征", "反应强度、峰值时间、恢复斜率、晚期残留"),
        ("定位", "作为副线：支持情绪变化与个体差异问题"),
    ]
    x = 0.75
    for i, (title, body) in enumerate(steps):
        add_card(slide, title, body, x + i * 2.45, 2.35, 2.05, 1.35, [BLUE, TEAL, GREEN, GOLD, RED][i])
    add_multiline(
        slide,
        [
            "主分析聚焦 recovery，避免把诱发阶段和恢复阶段混成同一个心理/生理过程。",
            "主观 arousal 或 valence 可用于解释个体差异，但不作为生理恢复的替代测量。",
        ],
        0.9,
        5.35,
        11.5,
        0.85,
        size=13,
        bullet=True,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[7])
    add_header(slide, "Next Tasks", "ACTION PLAN", 12)
    add_text(slide, "VAD 主线：收稿化", 0.85, 1.45, 2.5, 0.24, size=11, color=RED, bold=True)
    add_multiline(
        slide,
        [
            "核实 DREAMER Dominance 图示方向，固定主编码并保留反向敏感性。",
            "补充置换检验、bootstrap 和逐类删除敏感性分析。",
            "主分析使用名词映射，形容词映射作为敏感性分析。",
            "完成摘要、方法、结果、讨论和可复现性声明。",
        ],
        0.85,
        1.85,
        5.45,
        2.5,
        size=12.2,
        bullet=True,
    )
    add_text(slide, "情绪变化副线：补强与定位", 7.0, 1.45, 2.7, 0.24, size=11, color=BLUE, bold=True)
    add_multiline(
        slide,
        [
            "把 CASE 第二周结果整理为副线方法与发现摘要。",
            "明确副线只支持情绪变化/恢复动态，不抢占 VAD 主线。",
            "补充个体曲线与关键敏感性结果的展示材料。",
            "如果后续做人群实验，再单独决定任务范式和设备方案。",
        ],
        7.0,
        1.85,
        5.45,
        2.5,
        size=12.2,
        bullet=True,
    )
    add_text(slide, "执行原则：本月优先完成一条可投稿的 VAD 探索性论文；情绪变化方向作为副线，保留为后续研究和讨论支撑。", 0.85, 5.6, 11.4, 0.34, size=13, color=INK, bold=True)

    slide = prs.slides.add_slide(prs.slide_layouts[7])
    add_text(slide, "Thank you", 0.72, 0.78, 6.2, 0.62, size=34, color=BLUE, bold=True)
    add_rule(slide, 2.1)
    add_text(slide, "本周汇报主旨", 0.72, 3.03, 2.0, 0.25, size=10.5, color=MUTED, bold=True)
    add_text(
        slide,
        "VAD 方向已经形成可展示、可写作的探索性实证结果；\n情绪变化方向作为副线，为后续恢复动态与个体差异研究提供方法模板。",
        0.72,
        3.48,
        8.6,
        0.88,
        size=17,
        color=INK,
        bold=True,
    )
    add_text(slide, "Next: VAD 稳健性分析 + 情绪变化副线材料补强", 0.72, 5.6, 6.4, 0.28, size=12, color=BLUE, bold=True)
    add_footer(slide, 13)

    prs.save(str(OUTPUT))
    print("business review deck generated")


if __name__ == "__main__":
    build()
