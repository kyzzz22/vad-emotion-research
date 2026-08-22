"""Apply reviewer-facing revisions to the two-page Japanese research brief."""

from __future__ import annotations

import copy
import io
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(
    "/Users/mac/Downloads/研究室/EMB相关/"
    "研究概要書_JIANGXIJIE_MA26006_背景貢献強化版.docx"
)
FIGURE = (
    ROOT
    / "results"
    / "dreamer_nrc_robustness"
    / "figure_overview_main_readable.png"
)
OUTPUT = ROOT / "paper" / "研究概要書_JIANGXIJIE_MA26006_審査指摘反映版.docx"
AFFILIATION = (
    "所属　芝浦工業大学大学院 理工学研究科 "
    "修士課程 電気電子情報工学専攻"
)
AFFILIATION_IMAGE = ROOT / "paper" / "assets" / "affiliation_line.png"
UNICODE_FONT = "Arial Unicode MS"


def paragraph_starting(document: Document, prefix: str):
    matches = [p for p in document.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph starting with {prefix!r}, got {len(matches)}")
    return matches[0]


def remove_all_runs(paragraph) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def replace_plain(paragraph, text: str) -> None:
    first_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            first_rpr = copy.deepcopy(run._r.rPr)
            break
    remove_all_runs(paragraph)
    run = paragraph.add_run(text)
    if first_rpr is not None:
        run._r.insert(0, first_rpr)
    set_run_font(run, UNICODE_FONT)


def replace_inline_heading(paragraph, heading: str, body: str) -> None:
    heading_rpr = None
    body_rpr = None
    if paragraph.runs:
        if paragraph.runs[0]._r.rPr is not None:
            heading_rpr = copy.deepcopy(paragraph.runs[0]._r.rPr)
        if len(paragraph.runs) > 1 and paragraph.runs[1]._r.rPr is not None:
            body_rpr = copy.deepcopy(paragraph.runs[1]._r.rPr)
    remove_all_runs(paragraph)
    heading_run = paragraph.add_run(heading)
    if heading_rpr is not None:
        heading_run._r.insert(0, heading_rpr)
    set_run_font(heading_run, UNICODE_FONT)
    body_run = paragraph.add_run(body)
    if body_rpr is not None:
        body_run._r.insert(0, body_rpr)
    set_run_font(body_run, UNICODE_FONT)


def remove_paragraph(paragraph) -> None:
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)


def add_reference(document: Document, text: str) -> None:
    exemplar = paragraph_starting(document, "[6]")
    paragraph = document.add_paragraph()
    if exemplar._p.pPr is not None:
        paragraph._p.insert(0, copy.deepcopy(exemplar._p.pPr))
    run = paragraph.add_run(text)
    if exemplar.runs and exemplar.runs[0]._r.rPr is not None:
        run._r.insert(0, copy.deepcopy(exemplar.runs[0]._r.rPr))


def set_run_font(run, font_name: str) -> None:
    run.font.name = font_name
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = rpr._add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attribute}"), font_name)


def make_affiliation_image() -> None:
    font_path = Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf")
    if not font_path.exists():
        raise FileNotFoundError(font_path)
    font = ImageFont.truetype(str(font_path), 42)
    scratch = Image.new("RGB", (2400, 120), "white")
    draw = ImageDraw.Draw(scratch)
    box = draw.textbbox((0, 0), AFFILIATION, font=font)
    width = box[2] - box[0] + 28
    height = box[3] - box[1] + 18
    image = Image.new("RGB", (width, height), "white")
    ImageDraw.Draw(image).text((14, 7 - box[1]), AFFILIATION, fill="black", font=font)
    AFFILIATION_IMAGE.parent.mkdir(parents=True, exist_ok=True)
    image.save(AFFILIATION_IMAGE, dpi=(300, 300))


def make_text_image(text: str, filename: str) -> Path:
    font_path = Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf")
    font = ImageFont.truetype(str(font_path), 42)
    scratch = Image.new("RGB", (1000, 100), "white")
    draw = ImageDraw.Draw(scratch)
    box = draw.textbbox((0, 0), text, font=font, stroke_width=1)
    width = box[2] - box[0] + 12
    height = box[3] - box[1] + 12
    image = Image.new("RGB", (width, height), "white")
    ImageDraw.Draw(image).text(
        (6, 5 - box[1]),
        text,
        fill="black",
        font=font,
        stroke_width=1,
        stroke_fill="black",
    )
    path = ROOT / "paper" / "assets" / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, dpi=(300, 300))
    return path


def replace_with_text_image(paragraph, text: str, image_path: Path) -> None:
    remove_all_runs(paragraph)
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(image_path))
    inline_shape._inline.graphic.graphicData.pic.nvPicPr.cNvPr.set("descr", text)


def add_affiliation_line(metadata) -> None:
    paragraph = metadata.insert_paragraph_before()
    if metadata._p.pPr is not None:
        paragraph._p.insert(0, copy.deepcopy(metadata._p.pPr))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = 0
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(AFFILIATION_IMAGE), width=Inches(5.7))
    inline_shape._inline.graphic.graphicData.pic.nvPicPr.cNvPr.set(
        "descr", AFFILIATION
    )


def replace_embedded_figure(docx_path: Path, figure_path: Path) -> None:
    staged = docx_path.with_suffix(".staged.docx")
    with zipfile.ZipFile(docx_path, "r") as source_zip:
        with zipfile.ZipFile(staged, "w", compression=zipfile.ZIP_DEFLATED) as target_zip:
            for info in source_zip.infolist():
                payload = source_zip.read(info.filename)
                if info.filename.startswith("word/media/") and info.filename.endswith(
                    ".png"
                ):
                    with Image.open(io.BytesIO(payload)) as embedded:
                        # The scientific figure is the only tall raster; the
                        # affiliation line is a shallow text strip.
                        if embedded.height > 200:
                            payload = figure_path.read_bytes()
                target_zip.writestr(info, payload)
    staged.replace(docx_path)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not FIGURE.exists():
        raise FileNotFoundError(FIGURE)
    make_affiliation_image()
    heading_images = {
        "3 比較・評価手順": make_text_image(
            "3 比較・評価手順", "heading_3_comparison_procedure.png"
        ),
        "6 考察": make_text_image("6 考察", "heading_6_discussion.png"),
        "7 結論": make_text_image("7 結論", "heading_7_conclusion.png"),
        "8 今後の課題": make_text_image("8 今後の課題", "heading_8_future_work.png"),
    }

    document = Document(SOURCE)

    metadata = paragraph_starting(document, "指導教員")
    add_affiliation_line(metadata)
    replace_plain(
        metadata,
        "指導教員　菅谷みどり\tMA26006　姜晰頡",
    )

    replace_plain(
        paragraph_starting(document, "この問題を検証するには"),
        "この問題を検証するには，独立に整備されたデータ間の直接比較が必要である．"
        "NRC-VADは英語語彙のVAD規範辞書，DREAMERは映像視聴時のEEG・ECGと"
        "視聴後VAD評価，CASEは映像視聴時の生理信号と連続VA評価を収録した"
        "情動データセットである[2-4]．個々の妥当性だけでは，同じ感情カテゴリー名を"
        "介した対応，とくにVAD各次元の順位と値の広がりが保たれるかは判断できない．",
    )

    replace_plain(paragraph_starting(document, "3 提案する比較枠組み"), "3 比較・評価手順")
    replace_plain(
        paragraph_starting(document, "そこで，共通する感情カテゴリー名"),
        "共通する感情カテゴリー名を対応づけ，各尺度の理論上の端点を[-1，1]に"
        "線形変換する．この変換は尺度範囲をそろえるだけで，分布や測定特性の同一性を"
        "保証しない．そこで，探索的な構造比較として，①線形傾向，②順位，③値の広がりを"
        "分け，「関連」と「同値」を区別する．",
    )

    p41 = paragraph_starting(document, "4.1 評価目的")
    replace_inline_heading(
        p41,
        "4.1 評価目的　",
        "研究目的に対応させるため，本手順を9感情カテゴリーに適用する．"
        "主要な判断にはVADを含むDREAMERを用いる．一方，CASEはVAのみで共通カテゴリーも"
        "少ないため，処理手順の補足確認に限定する．",
    )

    replace_plain(
        paragraph_starting(document, "尺度範囲が異なるため"),
        "試行を9カテゴリーに集約し，線形傾向をPearson相関，順位をSpearman相関，"
        "値の広がりを回帰傾きで評価した．推論単位となる感情カテゴリー数が9と少ないため，"
        "9!通りの正確置換検定を用い，3次元のp値をHolm法で補正した．参加者と映像の抽出に"
        "よる不確実性は5,000回の階層bootstrapで評価した．Dominanceの符号反転感度分析も"
        "行った．",
    )

    remove_paragraph(paragraph_starting(document, "5 評価結果（続き）"))

    replace_plain(
        paragraph_starting(document, "線形対応が得られたVとA"),
        "線形対応が得られたVとAの回帰傾きはそれぞれ.520，.341で，いずれも1未満で"
        "あった（表1）．すなわち情動体験評価は中立点付近に縮小し，高相関でも数値は"
        "一致しない．ただし傾きは尺度，刺激，標本の影響を分離できない近似である．"
        "形容詞への置換でもV・Aの線形対応とDの判断は変わらず，Dominanceの符号反転でも"
        "「対応なし」の結論は変わらなかった．品詞差より各カテゴリーの極性・活性の大局的"
        "配置が結果を支配した可能性があるが，語彙数を増やした検証が必要である．CASEは"
        "4カテゴリーのため，主要結果の外的再現とはみなさない．",
    )

    replace_plain(
        paragraph_starting(document, "Dominanceでは線形対応も順位対応も"),
        "Dominanceでは線形対応も順位対応も得られず，共通の対応構造は支持されなかった．"
        "意味的評価が語に一般的に結びつく支配感を表すのに対し，情動体験は場面に対する"
        "自己・他者の責任／統制や状況統制の認知にも左右されうる[7]．評価対象の違いが"
        "対応を弱めた可能性があるが，本分析はこの機序を直接検証していない．",
    )

    replace_plain(
        paragraph_starting(document, "本分析は9カテゴリー"),
        "本分析は9カテゴリー，名詞1語，各2映像に限られ，辞書全体や多様な刺激には"
        "一般化できない．また，データセット間で参加者と課題が異なるため，個人内対応を"
        "判断できない．線形変換は尺度範囲をそろえるだけで，測定不変性も保証しない．"
        "さらに，CASEは4カテゴリーのVAのみであり，外的再現には不十分である．",
    )
    replace_plain(
        paragraph_starting(document, "そこで，まずDREAMERのDominance"),
        "まず，感情数・映像数・類義語を増やした外部データで再検証する．次に，"
        "同一参加者が感情語の意味と映像後体験を同一尺度で評価する実験を行い，刺激を"
        "交差させた混合効果モデルにより個人内対応を検証する．これにより，データセット差と"
        "尺度差を分離し，次元別較正の一般化可能性を評価する．",
    )

    for heading_text, image_path in heading_images.items():
        replace_with_text_image(
            paragraph_starting(document, heading_text),
            heading_text,
            image_path,
        )

    add_reference(
        document,
        '[7] Smith, C. A., Ellsworth, P. C. "Patterns of Cognitive Appraisal in Emotion". '
        "J. Pers. Soc. Psychol. 1985, Vol.48, No.4, p.813-838.",
    )

    if len(document.inline_shapes) != 6:
        raise ValueError(f"Expected six inline images, got {len(document.inline_shapes)}")
    shape = max(document.inline_shapes, key=lambda item: item.height)
    shape.height = int(shape.width * 920 / 2987)
    shape._inline.graphic.graphicData.pic.nvPicPr.cNvPr.set(
        "descr",
        "NRC-VADとDREAMERの9感情カテゴリーについて，"
        "Valence・Arousal・Dominance別に対応を示す散布図",
    )

    # Keep captions consistent with standard Japanese academic numbering.
    replace_plain(
        paragraph_starting(document, "図１"),
        "図1　意味的評価と情動体験評価の次元別対応（9感情カテゴリー）",
    )
    replace_plain(
        paragraph_starting(document, "表１"),
        "表1　VAD次元別の対応と尺度差（9感情カテゴリー）",
    )
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    replace_embedded_figure(OUTPUT, FIGURE)
    print(OUTPUT)


if __name__ == "__main__":
    main()
