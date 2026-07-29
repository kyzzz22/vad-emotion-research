from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper" / "DREAMER_NRC_完全分析報告_日本語.md"
OUTPUT = ROOT / "paper" / "DREAMER_NRC_完全分析報告_日本語.docx"

BODY_FONT = "Yu Mincho"
HEADING_FONT = "Yu Gothic"
MONO_FONT = "Consolas"
BLACK = "000000"
GRAY = "E6E6E6"
LIGHT_GRAY = "F2F2F2"
MID_GRAY = "B8B8B8"


def set_font(run, *, name=BODY_FONT, size=None, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run.font.color.rgb = RGBColor.from_string(BLACK)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=60, start=85, bottom=60, end=85) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_paragraph_border(paragraph, *, side="bottom", color=BLACK, size=8) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), "4")
    border.set(qn("w:color"), color)
    borders.append(border)


def set_paragraph_shading(paragraph, fill=LIGHT_GRAY) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, separate, end])


INLINE_PATTERN = re.compile(
    r"(\*\*.+?\*\*|`.+?`|\[[^]]+\]\([^)]+\)|(?<!\*)\*[^*\n]+?\*(?!\*))"
)


def add_inline(paragraph, text: str, *, size=None, default_font=BODY_FONT) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            set_font(paragraph.add_run(text[position : match.start()]), name=default_font, size=size)
        token = match.group(0)
        if token.startswith("**"):
            set_font(paragraph.add_run(token[2:-2]), name=default_font, size=size, bold=True)
        elif token.startswith("`"):
            set_font(paragraph.add_run(token[1:-1]), name=MONO_FONT, size=9)
        elif token.startswith("["):
            label = token[1 : token.index("]")]
            set_font(paragraph.add_run(label), name=default_font, size=size)
        else:
            set_font(paragraph.add_run(token[1:-1]), name=default_font, size=size, italic=True)
        position = match.end()
    if position < len(text):
        set_font(paragraph.add_run(text[position:]), name=default_font, size=size)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(3.5)
    normal.paragraph_format.widow_control = True

    for style_name, size, before, after in (
        ("Title", 24, 0, 12),
        ("Heading 1", 18, 12, 6),
        ("Heading 2", 14, 9, 4),
        ("Heading 3", 11.5, 6, 2.5),
    ):
        style = styles[style_name]
        style.font.name = HEADING_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    styles["Heading 1"].paragraph_format.page_break_before = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.left_indent = Cm(0.74)
        style.paragraph_format.first_line_indent = Cm(-0.37)
        style.paragraph_format.space_after = Pt(2)

    if "Report Caption" not in styles:
        caption = styles.add_style("Report Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = HEADING_FONT
        caption._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
        caption.font.size = Pt(9)
        caption.font.color.rgb = RGBColor.from_string(BLACK)
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(3)
        caption.paragraph_format.space_after = Pt(7)
        caption.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(
        header.add_run("DREAMER × NRC-VAD　完全分析報告書"),
        name=HEADING_FONT,
        size=8,
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("― "), name=HEADING_FONT, size=8)
    add_field(footer, "PAGE")
    set_font(footer.add_run(" ―"), name=HEADING_FONT, size=8)


def add_cover(doc: Document, lines: list[str]) -> int:
    title = lines[0][2:].strip()
    subtitle = lines[2][3:].strip()

    label = doc.add_paragraph()
    label.paragraph_format.space_after = Pt(54)
    set_font(
        label.add_run("研究分析報告書　2026年7月30日"),
        name=HEADING_FONT,
        size=9.5,
        bold=True,
    )

    p = doc.add_paragraph(style="Title")
    add_inline(p, title, size=24, default_font=HEADING_FONT)
    set_paragraph_border(p, side="bottom", color=BLACK, size=18)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(14)
    add_inline(p, subtitle, size=14, default_font=HEADING_FONT)
    for run in p.runs:
        run.bold = True

    index = 3
    while index < len(lines) and lines[index].strip() != "---":
        line = lines[index].strip()
        if line:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, line, size=10)
        index += 1

    return index + 1


def add_rule(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    set_paragraph_border(p, side="bottom", color=MID_GRAY, size=4)


def add_markdown_table(doc: Document, rows: list[str]) -> None:
    parsed = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in rows]
    if len(parsed) < 2:
        return
    data = [parsed[0]] + parsed[2:]
    columns = max(len(row) for row in data)
    table = doc.add_table(rows=len(data), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(data):
        prevent_row_split(table.rows[row_index])
        for column_index in range(columns):
            value = row[column_index] if column_index < len(row) else ""
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, GRAY if row_index == 0 else ("FAFAFA" if row_index % 2 == 0 else "FFFFFF"))
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value, size=8.2, default_font=HEADING_FONT)
            for run in paragraph.runs:
                set_font(
                    run,
                    name=HEADING_FONT,
                    size=8.2,
                    bold=(row_index == 0),
                )
    set_repeat_table_header(table.rows[0])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def add_figure(doc: Document, source_dir: Path, line: str) -> None:
    match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
    if not match:
        return
    caption, target = match.groups()
    path = (source_dir / target).resolve()
    if not path.exists():
        raise FileNotFoundError(f"Missing report figure: {path}")
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.add_run().add_picture(str(path), width=Cm(16.5))
    caption_paragraph = doc.add_paragraph(style="Report Caption")
    add_inline(caption_paragraph, caption, size=9, default_font=HEADING_FONT)


def add_quote(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.55)
    paragraph.paragraph_format.right_indent = Cm(0.35)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(7)
    set_paragraph_shading(paragraph)
    set_paragraph_border(paragraph, side="left", color=BLACK, size=12)
    add_inline(paragraph, text, size=10.5)


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.45)
    paragraph.paragraph_format.right_indent = Cm(0.45)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_together = True
    set_paragraph_shading(paragraph)
    set_paragraph_border(paragraph, side="left", color=BLACK, size=8)
    set_font(paragraph.add_run("\n".join(lines)), name=MONO_FONT, size=8.8)


def render_markdown(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    index = add_cover(doc, lines)
    in_executive_summary = True

    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            index += 1
            continue
        if line.startswith("```"):
            block: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            add_code_block(doc, block)
            index += 1
            continue
        if line.startswith("|"):
            table_rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_rows.append(lines[index].strip())
                index += 1
            add_markdown_table(doc, table_rows)
            continue
        if line.startswith("!["):
            add_figure(doc, SOURCE.parent, line)
            index += 1
            continue
        if line == "---":
            add_rule(doc)
            index += 1
            continue
        if line.startswith("# "):
            in_executive_summary = False
            paragraph = doc.add_paragraph(style="Heading 1")
            add_inline(paragraph, line[2:], size=18, default_font=HEADING_FONT)
            set_paragraph_border(paragraph, side="bottom", color=BLACK, size=10)
            index += 1
            continue
        if line.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 2")
            if in_executive_summary:
                paragraph.paragraph_format.page_break_before = False
            add_inline(paragraph, line[3:], size=14, default_font=HEADING_FONT)
            set_paragraph_border(paragraph, side="bottom", color=MID_GRAY, size=4)
            index += 1
            continue
        if line.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 3")
            add_inline(paragraph, line[4:], size=11.5, default_font=HEADING_FONT)
            index += 1
            continue
        if line.startswith("> "):
            add_quote(doc, line[2:])
            index += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            paragraph = doc.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\.\s+", "", line))
            index += 1
            continue
        if line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
            index += 1
            continue

        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
        add_inline(paragraph, line.rstrip())
        index += 1


def main() -> None:
    document = Document()
    configure_document(document)
    render_markdown(document, SOURCE.read_text(encoding="utf-8"))

    properties = document.core_properties
    properties.title = "語彙的VAD規範と映像誘発感情のクロスコンテクスト整合性"
    properties.subject = "DREAMER × NRC-VAD 完全分析報告書"
    properties.author = "VAD Emotion Research"
    properties.keywords = "VAD, DREAMER, NRC-VAD, 感情誘発, 語彙規範"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print("Japanese full report DOCX generated")


if __name__ == "__main__":
    main()
