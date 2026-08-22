from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / ".docx_template" / "AL23004_大橋華音_概要書v3.docx"
SOURCE = ROOT / "研究概要書_JIANGXIJIE_MA26006_修正版.docx"
UNWRAPPED = ROOT / ".docx_template" / "ohashi_template_unwrapped.docx"
INTERMEDIATE = ROOT / ".docx_template" / "研究概要書_大橋様式_作業中.docx"
OUTPUT = ROOT / "研究概要書_JIANGXIJIE_MA26006_大橋様式.docx"
PLOT = ROOT / ".docx_qa" / "image1_revised.png"

REFERENCE_SHA = "46fae160c511a1022f9f4f457310dd575abd42136f260b0c92876a8a3f6925c9"

MINCHO_NAME = "Hiragino Mincho ProN W3"
MINCHO_FILE = ROOT / ".docx_qa" / "HiraginoMinchoProN-W3.ttf"
GOTHIC_NAME = "Hiragino Sans W3"
GOTHIC_FILE = ROOT / ".docx_template" / "HiraginoKakuGothic-W3.ttf"

FONT_SPECS = [
    {
        "name": MINCHO_NAME,
        "alt": "ヒラギノ明朝 ProN W3",
        "file": MINCHO_FILE,
        "guid": "A27E2D4C-7F9B-4A5F-9C23-18D4A6B7E901",
        "rel_id": "rIdHiraginoMinchoFull",
        "target": "fonts/hiragino_mincho_full.odttf",
        "family": "roman",
    },
    {
        "name": GOTHIC_NAME,
        "alt": "ヒラギノ角ゴシック W3",
        "file": GOTHIC_FILE,
        "guid": "B6841A73-9D25-44D0-AEC1-70F8C345129B",
        "rel_id": "rIdHiraginoGothicFull",
        "target": "fonts/hiragino_gothic_full.odttf",
        "family": "swiss",
    },
]


def verify_reference():
    digest = sha256(REFERENCE.read_bytes()).hexdigest()
    if digest != REFERENCE_SHA:
        raise RuntimeError("Reference template changed; fresh distillation required.")


def unwrap_body_content_control():
    """Create a working copy with the reference's rich body content unwrapped."""
    temp = UNWRAPPED.with_suffix(".tmp.docx")
    with ZipFile(REFERENCE, "r") as src, ZipFile(
        temp, "w", compression=ZIP_DEFLATED
    ) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "word/document.xml":
                root = etree.fromstring(data)
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                body = root.find("w:body", namespaces=ns)
                for sdt in list(body.findall("w:sdt", namespaces=ns)):
                    content = sdt.find("w:sdtContent", namespaces=ns)
                    index = body.index(sdt)
                    body.remove(sdt)
                    if content is not None:
                        for child in list(content):
                            body.insert(index, child)
                            index += 1
                data = etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            dst.writestr(item, data)
    temp.replace(UNWRAPPED)


def clear_runs(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_run_font(run, east_asia, size, bold=False, latin=None):
    run.font.name = latin or east_asia
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), latin or east_asia)
    fonts.set(qn("w:hAnsi"), latin or east_asia)
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:cs"), latin or east_asia)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "ja-JP")


def format_paragraph(paragraph, role):
    pf = paragraph.paragraph_format
    pf.keep_together = False
    pf.widow_control = False
    if role == "title":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.space_before = Pt(0)
        pf.space_after = Pt(3.6)
        pf.line_spacing = Pt(14.5)
    elif role == "metadata":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.space_before = Pt(0)
        pf.space_after = Pt(3.6)
        pf.line_spacing = Pt(10.5)
        pf.tab_stops.add_tab_stop(Inches(6.45), WD_TAB_ALIGNMENT.RIGHT)
    elif role == "heading":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.space_before = Pt(3.6)
        pf.space_after = Pt(3.6)
        pf.line_spacing = Pt(10.8)
        pf.keep_with_next = True
        pf.first_line_indent = None
    elif role == "body":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.space_before = Pt(0)
        pf.space_after = Pt(3.6)
        pf.line_spacing = Pt(10.2)
        pf.first_line_indent = Pt(9)
    elif role == "caption":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.space_before = Pt(1.5)
        pf.space_after = Pt(3.6)
        pf.line_spacing = Pt(10.2)
        pf.keep_with_next = True
    elif role == "reference":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.space_before = Pt(0)
        pf.space_after = Pt(3.6)
        pf.line_spacing = Pt(10)
        pf.first_line_indent = None


def set_text(paragraph, text, role):
    clear_runs(paragraph)
    run = paragraph.add_run(text)
    if role == "title":
        set_run_font(run, GOTHIC_NAME, 14, bold=True)
    elif role == "metadata":
        set_run_font(run, GOTHIC_NAME, 10, bold=False)
    elif role == "heading":
        set_run_font(run, GOTHIC_NAME, 10.5, bold=False)
    elif role == "caption":
        set_run_font(run, MINCHO_NAME, 9, bold=False)
    elif role == "reference":
        set_run_font(run, MINCHO_NAME, 9, bold=False, latin="Times New Roman")
    else:
        set_run_font(run, MINCHO_NAME, 9, bold=False, latin="Times New Roman")
    format_paragraph(paragraph, role)


def set_columns(section, count):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    for child in list(cols):
        cols.remove(child)
    if count == 1:
        cols.set(qn("w:num"), "1")
        cols.set(qn("w:space"), "720")
        cols.attrib.pop(qn("w:equalWidth"), None)
    else:
        cols.set(qn("w:num"), "2")
        cols.set(qn("w:space"), "298")
        cols.set(qn("w:equalWidth"), "0")
        first = OxmlElement("w:col")
        first.set(qn("w:w"), "4499")
        first.set(qn("w:space"), "298")
        second = OxmlElement("w:col")
        second.set(qn("w:w"), "4499")
        second.set(qn("w:space"), "0")
        cols.extend([first, second])


def strip_body_after_metadata(doc):
    body = doc._element.body
    metadata = doc.paragraphs[2]._p
    keep_index = body.index(metadata)
    for child in list(body)[keep_index + 1 :]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_heading(doc, text):
    paragraph = doc.add_paragraph()
    set_text(paragraph, text, "heading")
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    set_text(paragraph, text, "body")
    return paragraph


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    set_text(paragraph, text, "caption")
    return paragraph


def remove_prefix(text):
    if "　" in text:
        return text.split("　", 1)[1]
    return text


def apply_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge, val, size, color in [
        ("top", "single", "8", "000000"),
        ("bottom", "single", "8", "000000"),
        ("insideH", "single", "4", "BFBFBF"),
        ("left", "nil", "0", "FFFFFF"),
        ("right", "nil", "0", "FFFFFF"),
        ("insideV", "nil", "0", "FFFFFF"),
    ]:
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), val)
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in [("top", 35), ("left", 55), ("bottom", 35), ("right", 55)]:
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_results_table(doc, source_table):
    rows = [[cell.text for cell in row.cells] for row in source_table.rows]
    table = doc.add_table(rows=len(rows), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [2200, 2366, 2366, 2366]
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9298")
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    apply_table_borders(table)
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[c_idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(9.6)
            clear_runs(paragraph)
            run = paragraph.add_run(rows[r_idx][c_idx])
            set_run_font(
                run,
                GOTHIC_NAME if r_idx == 0 else MINCHO_NAME,
                8.7 if r_idx else 9,
                bold=(r_idx == 0),
                latin=None if r_idx == 0 else "Times New Roman",
            )
    return table


def obfuscate_font(font_path, guid):
    key = bytes.fromhex(guid.replace("-", ""))[::-1]
    data = bytearray(font_path.read_bytes())
    for index in range(min(32, len(data))):
        data[index] ^= key[index % 16]
    return bytes(data)


def embed_fonts(input_docx, output_docx):
    font_blobs = {spec["target"]: obfuscate_font(spec["file"], spec["guid"]) for spec in FONT_SPECS}
    temp = output_docx.with_suffix(".tmp.docx")
    with ZipFile(input_docx, "r") as src, ZipFile(
        temp, "w", compression=ZIP_DEFLATED
    ) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "word/fontTable.xml":
                root = etree.fromstring(data)
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                for spec in FONT_SPECS:
                    found = root.xpath(
                        f".//w:font[@w:name='{spec['name']}']", namespaces=ns
                    )
                    font = found[0] if found else etree.SubElement(root, qn("w:font"))
                    font.set(qn("w:name"), spec["name"])
                    for child in list(font):
                        if etree.QName(child).localname.startswith("embed"):
                            font.remove(child)
                    alt = etree.SubElement(font, qn("w:altName"))
                    alt.set(qn("w:val"), spec["alt"])
                    family = etree.SubElement(font, qn("w:family"))
                    family.set(qn("w:val"), spec["family"])
                    pitch = etree.SubElement(font, qn("w:pitch"))
                    pitch.set(qn("w:val"), "variable")
                    embedded = etree.SubElement(font, qn("w:embedRegular"))
                    embedded.set(qn("r:id"), spec["rel_id"])
                    embedded.set(qn("w:fontKey"), "{" + spec["guid"] + "}")
                data = etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            elif item.filename == "word/_rels/fontTable.xml.rels":
                root = etree.fromstring(data)
                rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
                ids = {spec["rel_id"] for spec in FONT_SPECS}
                for rel in list(root):
                    if rel.get("Id") in ids:
                        root.remove(rel)
                for spec in FONT_SPECS:
                    rel = etree.SubElement(root, f"{{{rel_ns}}}Relationship")
                    rel.set("Id", spec["rel_id"])
                    rel.set(
                        "Type",
                        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font",
                    )
                    rel.set("Target", spec["target"])
                data = etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            elif item.filename == "[Content_Types].xml":
                root = etree.fromstring(data)
                ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
                targets = {"/word/" + spec["target"] for spec in FONT_SPECS}
                for override in list(root):
                    if override.get("PartName") in targets:
                        root.remove(override)
                for target in sorted(targets):
                    override = etree.SubElement(root, f"{{{ct_ns}}}Override")
                    override.set("PartName", target)
                    override.set(
                        "ContentType",
                        "application/vnd.openxmlformats-officedocument.obfuscatedFont",
                    )
                data = etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            dst.writestr(item, data)
        for target, blob in font_blobs.items():
            dst.writestr("word/" + target, blob)
    temp.replace(output_docx)


def main():
    verify_reference()
    unwrap_body_content_control()
    source = Document(SOURCE)
    doc = Document(UNWRAPPED)

    # Retain the reference title/metadata slots and their section transition.
    set_text(
        doc.paragraphs[0],
        "感情語VADの誘発情動体験へのクロスコンテクスト移送可能性の検討",
        "title",
    )
    blank = doc.paragraphs[1]
    clear_runs(blank)
    blank.paragraph_format.space_before = Pt(0)
    blank.paragraph_format.space_after = Pt(0)
    blank.paragraph_format.line_spacing = Pt(1)
    set_text(
        doc.paragraphs[2],
        "指導教員　菅谷みどり\tMA26006　姜晰頡",
        "metadata",
    )
    strip_body_after_metadata(doc)
    set_columns(doc.sections[-1], 2)

    p = source.paragraphs
    add_heading(doc, "1 背景・課題")
    for index in [4, 5, 6, 7]:
        add_body(doc, p[index].text)

    add_heading(doc, "2 目的・提案")
    add_body(doc, p[9].text)
    add_body(doc, p[11].text)

    add_heading(doc, "3 二次データ分析")
    add_heading(doc, "3.1 分析方針")
    add_body(doc, remove_prefix(p[13].text))
    add_heading(doc, "3.2 対象データ")
    add_body(doc, remove_prefix(p[14].text))
    add_heading(doc, "3.3 分析手順")
    add_body(doc, p[16].text)
    add_heading(doc, "3.4 分析結果")
    add_body(doc, p[18].text)

    figure_section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(figure_section, 1)
    figure_paragraph = doc.add_paragraph()
    figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_paragraph.paragraph_format.space_before = Pt(0)
    figure_paragraph.paragraph_format.space_after = Pt(0)
    figure_paragraph.add_run().add_picture(str(PLOT), width=Inches(6.25))
    add_caption(
        doc,
        "図1 : 意味的評価と情動体験評価のクロスコンテクスト対応（9感情カテゴリー）",
    )

    page_break = doc.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)
    add_caption(doc, "表1 : VAD次元別の対応と尺度差（9感情カテゴリー）")
    add_results_table(doc, source.tables[0])

    remainder_section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(remainder_section, 2)
    add_heading(doc, "3.4 分析結果（続き）")
    add_body(doc, p[26].text)
    add_heading(doc, "3.5 考察")
    for index in [28, 29, 30]:
        add_body(doc, p[index].text)

    add_heading(doc, "4 まとめと展望")
    for index in [32, 34, 35]:
        add_body(doc, p[index].text)

    add_heading(doc, "参考文献")
    for index in range(37, 43):
        paragraph = doc.add_paragraph()
        set_text(paragraph, p[index].text, "reference")

    doc.save(INTERMEDIATE)
    embed_fonts(INTERMEDIATE, OUTPUT)


if __name__ == "__main__":
    main()
